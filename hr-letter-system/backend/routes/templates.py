from flask import Blueprint, request, jsonify, current_app, send_file
from flask_jwt_extended import jwt_required, get_jwt_identity
from datetime import datetime
from bson import ObjectId
from io import BytesIO
import os, re, uuid, gridfs
from docx import Document

templates_bp = Blueprint('templates', __name__)

DELETE_ROLES = {'admin', 'hr_head'}


def get_fs():
    return gridfs.GridFS(current_app.db, collection='templates_fs')


def extract_placeholders_from_bytes(data):
    doc  = Document(BytesIO(data))
    text = '\n'.join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += '\n' + cell.text
    return sorted(set(re.findall(r'\{\{(\w+)\}\}', text)))


def serialize(t):
    t['_id'] = str(t['_id'])
    return t


def _get_caller(db, uid):
    user = db.users.find_one({'_id': ObjectId(uid)})
    if not user:
        return None, (jsonify({'error': 'User not found'}), 404)
    return user, None


@templates_bp.route('/', methods=['GET'])
@jwt_required()
def list_templates():
    db        = current_app.db
    tmpl_type = request.args.get('type')
    query     = {'type': tmpl_type} if tmpl_type else {}
    tmpl      = list(db.templates.find(query).sort('created_at', -1))
    return jsonify([serialize(t) for t in tmpl])


@templates_bp.route('/<tid>', methods=['GET'])
@jwt_required()
def get_template(tid):
    db = current_app.db
    t  = db.templates.find_one({'_id': ObjectId(tid)})
    return jsonify(serialize(t)) if t else (jsonify({'error': 'Not found'}), 404)


@templates_bp.route('/upload', methods=['POST'])
@jwt_required()
def upload():
    db   = current_app.db
    uid  = get_jwt_identity()
    file = request.files.get('file')
    if not file:
        return jsonify({'error': 'No file provided'}), 400

    name      = request.form.get('name') or file.filename
    tmpl_type = request.form.get('type', 'offer')
    if tmpl_type not in ('offer', 'relieving', 'appointment_order'):
        return jsonify({'error': "type must be 'offer', 'relieving', or 'appointment_order'"}), 400

    latest  = db.templates.find_one({'type': tmpl_type}, sort=[('version', -1)])
    version = (latest['version'] + 1) if latest else 1

    slug     = uuid.uuid4().hex[:8]
    filename = f"{tmpl_type}_v{version}_{slug}_{file.filename}"

    # Read file bytes once
    data = file.read()

    try:
        placeholders = extract_placeholders_from_bytes(data)

        # Save to GridFS
        fs = get_fs()
        gridfs_id = str(fs.put(data, filename=filename, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'))

        doc = {
            'name':         name,
            'type':         tmpl_type,
            'version':      version,
            'placeholders': placeholders,
            'gridfs_id':    gridfs_id,
            'filename':     filename,
            'is_active':    True,
            'uploaded_by':  uid,
            'created_at':   datetime.utcnow(),
        }
        result = db.templates.insert_one(doc)
    except Exception as e:
        return jsonify({'error': f'Template processing failed: {str(e)}'}), 500

    return jsonify({'id': str(result.inserted_id), 'placeholders': placeholders, 'version': version}), 201


@templates_bp.route('/<tid>/toggle', methods=['POST'])
@jwt_required()
def toggle(tid):
    db  = current_app.db
    uid = get_jwt_identity()
    caller, err = _get_caller(db, uid)
    if err: return err
    if caller.get('role') != 'admin':
        return jsonify({'error': 'Admin only'}), 403
    t = db.templates.find_one({'_id': ObjectId(tid)})
    if not t:
        return jsonify({'error': 'Not found'}), 404
    new_val = not t['is_active']
    db.templates.update_one({'_id': ObjectId(tid)}, {'$set': {'is_active': new_val}})
    return jsonify({'is_active': new_val})


@templates_bp.route('/<tid>/download', methods=['GET'])
@jwt_required()
def download(tid):
    db = current_app.db
    t  = db.templates.find_one({'_id': ObjectId(tid)})
    if not t:
        return jsonify({'error': 'Not found'}), 404

    # Try GridFS first, fall back to disk for legacy templates
    if t.get('gridfs_id'):
        fs = get_fs()
        grid_out = fs.get(ObjectId(t['gridfs_id']))
        return send_file(
            BytesIO(grid_out.read()),
            as_attachment=True,
            download_name=t['filename'],
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
    elif t.get('file_path') and os.path.exists(t['file_path']):
        return send_file(t['file_path'], as_attachment=True, download_name=t['filename'])
    else:
        return jsonify({'error': 'Template file not found. Please re-upload.'}), 404


@templates_bp.route('/<tid>', methods=['DELETE'])
@jwt_required()
def delete_template(tid):
    db  = current_app.db
    uid = get_jwt_identity()

    caller, err = _get_caller(db, uid)
    if err: return err

    if caller.get('role') not in DELETE_ROLES:
        return jsonify({'error': 'Only Admin or HR Head can delete templates'}), 403

    t = db.templates.find_one({'_id': ObjectId(tid)})
    if not t:
        return jsonify({'error': 'Template not found'}), 404

    linked_count = db.letters.count_documents({'template_id': tid})
    if linked_count > 0:
        return jsonify({
            'error': (
                f"Cannot delete: {linked_count} offer letter(s) reference this template. "
                "Deactivate it instead to prevent new use while preserving the audit trail."
            )
        }), 400

    # Delete from GridFS
    if t.get('gridfs_id'):
        try:
            fs = get_fs()
            fs.delete(ObjectId(t['gridfs_id']))
        except Exception:
            pass

    # Delete from disk (legacy)
    fpath = t.get('file_path')
    if fpath and os.path.exists(fpath):
        try:
            os.remove(fpath)
        except OSError:
            pass

    db.templates.delete_one({'_id': ObjectId(tid)})
    return jsonify({
        'message': f"Template '{t['name']}' deleted by {caller.get('name', uid)} ({caller.get('role')}) at {datetime.utcnow().isoformat()}"
    })